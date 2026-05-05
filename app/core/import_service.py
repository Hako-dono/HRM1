from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from data.import_repository import ImportRepository, RecipientRecord

REQUIRED_PAYROLL_COLUMNS = [
    "Mã NV", "Họ và tên", "Chức vụ", "Phòng ban", "Lương cơ bản", "Ngày công chuẩn", "Ngày công thực tế", "Thực nhận"
]
REQUIRED_EMPLOYEE_COLUMNS = ["Mã NV", "Email", "Tình trạng"]

OPTIONAL_PAYROLL_COLUMNS = [
    "Phụ cấp ăn trưa", "Phụ cấp trang phục", "Phụ cấp điện thoại", "Lương hiệu suất", "Lương làm thêm giờ",
    "Đóng góp trách nhiệm", "BHXH", "BHYT", "BHTN", "Thuế TNCN", "Tạm ứng"
]


@dataclass
class ImportResult:
    session_id: int
    payroll_month: str
    warnings: list[str]


class ImportErrorMessage(Exception):
    pass


class ImportService:
    def __init__(self, repository: ImportRepository) -> None:
        self.repository = repository

    def import_files(self, payroll_path: Path, employee_path: Path, manual_month: str | None = None) -> ImportResult:
        payroll_rows, payroll_month = self._read_payroll(payroll_path)
        if not payroll_month:
            payroll_month = (manual_month or "").strip()
            if not re.match(r"^\d{4}-\d{2}$", payroll_month):
                raise ImportErrorMessage("Không parse được tháng lương. Vui lòng nhập theo định dạng YYYY-MM.")

        employee_map, duplicate_employees = self._read_employees(employee_path)
        recipients: list[RecipientRecord] = []
        warnings: list[str] = []

        for code, row in payroll_rows.items():
            emp = employee_map.get(code)
            email = emp["Email"] if emp else ""
            status = emp["Tình trạng"] if emp else ""
            validation_status = "valid"
            messages: list[str] = []
            sendable = True

            if code in duplicate_employees:
                validation_status = "warning"
                messages.append("Mã NV bị trùng trong file nhân sự")

            net_amount = self._to_number(row.get("Thực nhận"))
            if net_amount is None:
                validation_status = "error"
                messages.append("Thực nhận trống hoặc không hợp lệ")
                sendable = False

            if not email:
                validation_status = "warning" if validation_status != "error" else validation_status
                messages.append("Thiếu email")
                sendable = False
            elif not self._is_valid_email(email):
                validation_status = "error"
                messages.append("Email không hợp lệ")
                sendable = False

            payroll_data = {}
            for col in REQUIRED_PAYROLL_COLUMNS + OPTIONAL_PAYROLL_COLUMNS:
                value = row.get(col)
                if col in OPTIONAL_PAYROLL_COLUMNS and (value is None or str(value).strip() == ""):
                    value = 0
                payroll_data[col] = value

            recipients.append(
                RecipientRecord(
                    employee_code=code,
                    employee_name=str(row.get("Họ và tên", "") or ""),
                    job_title=str(row.get("Chức vụ", "") or ""),
                    department=str(row.get("Phòng ban", "") or ""),
                    employee_email=email,
                    employment_status=status,
                    net_amount=net_amount,
                    payroll_data=payroll_data,
                    validation_status=validation_status,
                    validation_message="; ".join(messages),
                    sendable=sendable,
                )
            )

        if duplicate_employees:
            warnings.append(f"Có {len(duplicate_employees)} Mã NV bị trùng trong file nhân sự.")

        session_id = self.repository.create_or_replace_draft_session(payroll_month, payroll_path, employee_path, recipients)
        return ImportResult(session_id=session_id, payroll_month=payroll_month, warnings=warnings)

    def _read_payroll(self, path: Path) -> tuple[dict[str, dict], str | None]:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        payroll_month = self._extract_month_from_sheet(ws)

        header_row_idx = None
        headers = []
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=50, values_only=True), start=1):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if "Mã NV" in cells and "Thực nhận" in cells:
                header_row_idx = i
                headers = cells
                break
        if header_row_idx is None:
            raise ImportErrorMessage("Không tìm thấy header payroll hợp lệ.")

        missing = [c for c in REQUIRED_PAYROLL_COLUMNS if c not in headers]
        if missing:
            raise ImportErrorMessage(f"Thiếu cột bắt buộc payroll: {', '.join(missing)}")

        code_idx = headers.index("Mã NV")
        rows: dict[str, dict] = {}
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            data = {headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))}
            code = str(data.get("Mã NV", "") or "").strip()
            if not code:
                continue
            if code in rows:
                raise ImportErrorMessage(f"File payroll có Mã NV trùng: {code}")
            rows[code] = data
        return rows, payroll_month

    def _read_employees(self, path: Path) -> tuple[dict[str, dict], set[str]]:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        headers = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        missing = [c for c in REQUIRED_EMPLOYEE_COLUMNS if c not in headers]
        if missing:
            raise ImportErrorMessage(f"Thiếu cột bắt buộc employee: {', '.join(missing)}")

        data = {}
        duplicates: set[str] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            item = {headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))}
            code = str(item.get("Mã NV", "") or "").strip()
            if not code:
                continue
            if code in data:
                duplicates.add(code)
            data[code] = {
                "Email": str(item.get("Email", "") or "").strip(),
                "Tình trạng": str(item.get("Tình trạng", "") or "").strip(),
            }
        return data, duplicates

    def _extract_month_from_sheet(self, ws) -> str | None:
        for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            for cell in row:
                if not cell:
                    continue
                txt = str(cell)
                m = re.search(r"(20\d{2})\D{0,3}(0?[1-9]|1[0-2])", txt)
                if m:
                    return f"{m.group(1)}-{int(m.group(2)):02d}"
        return None

    def _to_number(self, value) -> float | None:
        if value is None:
            return None
        try:
            text = str(value).replace(",", "").strip()
            return float(text)
        except Exception:
            return None

    def _is_valid_email(self, email: str) -> bool:
        return re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email) is not None
